from pathlib import Path
from docx import Document

MAX_DOCX_SIZE_BYTES = 40 * 1024 * 1024 # 40 MB

def extract_text_from_docx(file_path: str | Path) -> str:
    """Extract text from paragraphs and tables in a DOCX document."""

    path = Path(file_path).expanduser()

    if path.is_symlink():
        raise PermissionError("Symbolic-link files are not allowed.")

    try:
        resolved_path = path.resolve(strict=True)

        if not resolved_path.is_file():
            raise ValueError("The provided path is not a file.")

        if resolved_path.suffix.lower() != ".docx":
            raise ValueError("The file must have a .docx extension.")

        if resolved_path.stat().st_size > MAX_DOCX_SIZE_BYTES:
            raise ValueError("The DOCX file exceeds the 40 MB size limit.")

        document = Document(str(resolved_path))

        extracted_parts = []

        # Extract paragraphs.
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                extracted_parts.append(text)

        # Extract table contents row by row.
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]
                if any(cells):
                    extracted_parts.append(" | ".join(cells))

        return "\n".join(extracted_parts)

    except (OSError, Exception) as exc:
        if isinstance(exc, (ValueError, PermissionError)):
            raise

        raise ValueError(
            "The DOCX document could not be opened or processed."
        ) from exc