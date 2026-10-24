from pathlib import Path
from typing import Any
from openpyxl import load_workbook

MAX_EXCEL_SIZE_BYTES = 100 * 1024 * 1024  # 100 MB
MAX_WORKSHEETS = 100
MAX_ROWS_PER_SHEET = 10000
MAX_COLUMNS_PER_SHEET = 500

def extract_text_from_excel(
    file_path: str | Path,
) -> str:
    """Extract worksheet names and cell values from an XLSX workbook."""

    path = Path(file_path).expanduser()

    if path.is_symlink():
        raise PermissionError("Symbolic-link files are not allowed.")

    try:
        resolved_path = path.resolve(strict=True)

        if not resolved_path.is_file():
            raise ValueError("The provided path is not a file.")

        if resolved_path.suffix.lower() != ".xlsx":
            raise ValueError("The file must have an .xlsx extension.")

        if resolved_path.stat().st_size > MAX_EXCEL_SIZE_BYTES:
            raise ValueError("The workbook exceeds the 100 MB size limit.")

        workbook = load_workbook(
            filename=resolved_path,
            read_only=True,
            data_only=True,
        )

    except (OSError, ValueError) as exc:
        raise ValueError(
            f"Could not open Excel workbook: {exc}"
        ) from exc

    extracted_parts: list[str] = []

    try:
        if len(workbook.sheetnames) > MAX_WORKSHEETS:
            raise ValueError(
                f"The workbook exceeds {MAX_WORKSHEETS} worksheets."
            )

        for worksheet in workbook.worksheets:
            extracted_parts.append(
                f"Worksheet: {worksheet.title}"
            )

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    max_row=MAX_ROWS_PER_SHEET,
                    max_col=MAX_COLUMNS_PER_SHEET,
                    values_only=True,
                ),
                start=1,
            ):  
                # Find the last cell that contains a value.
                last_non_empty = next(
                    (
                        index
                        for index in range(len(row) - 1, -1, -1)
                        if row[index] is not None
                        and str(row[index]).strip()
                    ),
                    -1,
                )

                # Skip rows containing no actual data.
                if last_non_empty == -1:
                    continue

                values = [
                    str(value).strip() if value is not None else ""
                    for value in row[: last_non_empty + 1]
                ]

                # Skip entirely empty rows.
                if not any(values):
                    continue

                extracted_parts.append(
                    f"Row {row_number}: " + " | ".join(values)
                )

    finally:
        workbook.close()

    return "\n".join(extracted_parts)
